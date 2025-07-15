"""
Модуль для генерации документов DOCX с актами списания
"""
import io
import os
import httpx
import logging
from pathlib import Path
from datetime import datetime, date
from typing import Dict, List, Optional
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from fastapi import BackgroundTasks
from sqlalchemy.ext.asyncio import AsyncSession
from services.write_off import WriteOffService
from models.group import Group
from sqlalchemy.future import select

# Настройка логирования
logger = logging.getLogger(__name__)

# --- ДОБАВЛЯЕМ СЛОВАРЬ ПЕРЕВОДА ПРИЧИН ---
REASON_TRANSLATE = {
    "consumption": "Расход",
    "spoilage": "Порча",
    "loss": "Потеря",
    "marketing": "Маркетинг",
    "client_refusal": "Отказ клиента",
    "defect": "Брак",
    "salary_withholding": "Удержание из з/п",
    "testing": "Проработка",
    # ... дополни по необходимости
}
# --- КОНЕЦ СЛОВАРЯ ---

def create_write_off_document(data: Dict) -> io.BytesIO:
    """
    Создает DOCX документ акта списания на основе переданных данных
    
    Args:
        data: словарь с данными для акта списания
            - chatTitle: название чата/филиала
            - date: дата акта списания
            - items: список элементов списания
    
    Returns:
        Документ в формате BytesIO
    """
    # Создаем документ
    doc = Document()
    
    # Устанавливаем поля страницы
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    # Добавляем заголовок документа
    title = doc.add_paragraph("Акт списания")
    title_format = title.paragraph_format
    title_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format.space_after = Pt(12)
    title_run = title.runs[0]
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    
    # Форматирование даты и филиала
    chat_title = data.get('chatTitle', 'Неизвестный филиал')
    
    # Форматируем дату из ISO в читаемый формат
    date_str = data.get('date', datetime.now().isoformat())
    try:
        if isinstance(date_str, str):
            date_obj = datetime.fromisoformat(date_str.replace('Z', '+00:00'))
            date_formatted = date_obj.strftime("%d.%m.%Y %H:%M")
        else:
            date_formatted = datetime.now().strftime("%d.%m.%Y %H:%M")
    except (ValueError, TypeError):
        date_formatted = datetime.now().strftime("%d.%m.%Y %H:%M")
    
    # Добавляем строку с датой и филиалом
    date_line = doc.add_paragraph(f"От {date_formatted}. Филиал: {chat_title}.")
    date_line.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.paragraph_format.space_after = Pt(12)
    
    # Создаем таблицу для элементов списания
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Устанавливаем ширину столбцов
    table.columns[0].width = Cm(6)  # Наименование
    table.columns[1].width = Cm(3)  # Кол-во
    table.columns[2].width = Cm(6)  # Причина списания
    
    # Заголовки таблицы
    header_cells = table.rows[0].cells
    header_cells[0].text = "Наименование"
    header_cells[1].text = "Кол-во"
    header_cells[2].text = "Причина списания"
    
    # Стиль заголовка таблицы
    for cell in header_cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    
    # Добавляем строки с элементами списания
    items = data.get('items', [])
    for item in items:
        row_cells = table.add_row().cells
        
        # Наименование
        row_cells[0].text = item.get('name', '')
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Количество с единицей измерения
        quantity = item.get('quantity', 0)
        unit_type = item.get('unitType', 'шт')
        row_cells[1].text = f"{quantity} {unit_type}"
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Причина списания
        reason = item.get('reason', {})
        reason_title = reason.get('title', '') if reason else ''
        # --- ПЕРЕВОДИМ НА РУССКИЙ ---
        reason_title = REASON_TRANSLATE.get(reason_title, reason_title)
        # --- КОНЕЦ ПЕРЕВОДА ---
        
        # Добавляем детальное описание в скобках, если оно существует
        description = item.get('description', '')
        if description:
            reason_text = f"{reason_title} ({description})"
        else:
            reason_text = reason_title
            
        row_cells[2].text = reason_text
        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Добавляем информацию о времени формирования документа
    doc.add_paragraph("")
    generation_info = doc.add_paragraph(f"Документ сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}")
    generation_info.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    generation_info.paragraph_format.space_after = Pt(12)
    
    
    # Добавляем раздел с подписями
    doc.add_paragraph("")
    doc.add_paragraph("")
    # --- ПОДСТАВЛЯЕМ ОТВЕТСТВЕННОГО ---
    responsible = data.get('responsible', '')
    # Сначала строка с ФИО, ниже — линия для подписи
    if responsible:
        fio_par = doc.add_paragraph(f"Ответственный:    {responsible}")
        fio_par.paragraph_format.space_after = Pt(2)
        line_par = doc.add_paragraph(" " * 29 + "      _____________________/ ________________")
        line_par.paragraph_format.space_after = Pt(12)
    else:
        signatures = doc.add_paragraph("Ответственный: _____________________/ ________________")
        signatures.paragraph_format.space_after = Pt(12)
    # --- КОНЕЦ ПОДСТАНОВКИ ---
    doc.add_paragraph("Администратор: ________________________ / _________________")
    doc.add_paragraph("")
    doc.add_paragraph("Руководитель: ________________________ / _________________")
    
    # Сохраняем документ в памяти
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return file_stream

# --- Новая бизнес-логика ---

async def generate_and_send_write_off_report(
    group_id: int,
    db: AsyncSession,
    background_tasks: BackgroundTasks,
    responsible_first_name: str = None,
    responsible_last_name: str = None,
    date_filter: Optional[date] = None
):
    """
    Генерирует DOCX-акт списания по группе, собирает фотографии, сохраняет файл и отправляет всё боту через background-задачу.
    
    Args:
        group_id: ID группы
        db: Сессия БД
        background_tasks: Фоновые задачи
        responsible_first_name: Имя ответственного
        responsible_last_name: Фамилия ответственного
        date_filter: Фильтр по дате (если None - все записи)
    """
    # 1. Получаем списания по группе с фильтром по дате
    service = WriteOffService()
    write_offs = await service.get_write_offs_by_group(db, group_id, date_filter)
    
    # Логируем информацию о фильтрации
    if date_filter:
        logger.info(f"📅 Генерация отчета для группы {group_id} за дату: {date_filter}")
    else:
        logger.info(f"📅 Генерация отчета для группы {group_id} за все даты")
    
    logger.info(f"📋 Найдено {len(write_offs)} списаний для отчета")

    # 2. Получаем название группы по group_id (Telegram chat_id)
    result = await db.execute(select(Group).where(Group.group_id == int(group_id)))
    group = result.scalar_one_or_none()
    chat_title = group.title if group else f"Группа {group_id}"

    # 3. Формируем структуру для генератора и собираем фотографии
    items = []
    photos = []  # Массив фотографий для отправки в бот
    
    for w in write_offs:
        # Добавляем данные для DOCX
        items.append({
            "name": w.name,
            "quantity": w.quantity,
            "unitType": getattr(w, "unit_type", "шт"),
            "reason": {"title": w.reason} if w.reason else {},
            "description": w.description or ""
        })
        
        # Собираем фотографии если есть
        if hasattr(w, 'photo_path') and w.photo_path:
            photo_full_path = f"/app/shared/write_off_photos/{w.photo_path}"
            if os.path.exists(photo_full_path):
                logger.info(f"📷 Найдено фото для {w.name}: {w.photo_path}")
                photos.append({
                    "file_path": photo_full_path,
                    "caption": f"📦 {w.name}\n💯 {w.quantity} {getattr(w, 'unit_type', 'шт')}\n🔹 {REASON_TRANSLATE.get(w.reason, w.reason) if w.reason else 'Без причины'}",
                    "item_name": w.name,
                    "quantity": w.quantity,
                    "unit_type": getattr(w, "unit_type", "шт"),
                    "reason": w.reason
                })
            else:
                logger.warning(f"❌ Фото не найдено: {photo_full_path}")
        else:
            logger.info(f"ℹ️ Нет фото для {w.name}")
    
    # --- СОБИРАЕМ ФИО ОТВЕТСТВЕННОГО ---
    responsible = ""
    if responsible_first_name or responsible_last_name:
        responsible = f"{responsible_first_name or ''} {responsible_last_name or ''}".strip()
    else:
        responsible = '________________________ / _________________'
    # --- КОНЕЦ СОБИРАНИЯ ФИО ---
    
    # Используем дату фильтра или текущую дату для отчета
    report_date = date_filter if date_filter else datetime.now().date()
    
    data = {
        "chatTitle": chat_title,
        "date": report_date.isoformat(),
        "items": items,
        "responsible": responsible
    }

    # 4. Генерируем DOCX
    file_stream = create_write_off_document(data)

    # 5. Сохраняем файл во временную папку
    reports_dir = Path(os.getenv("SHARED_WRITE_OFF_REPORTS", "/app/shared/write_off_reports"))
    reports_dir.mkdir(parents=True, exist_ok=True)
    filename = f"write_off_{group_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    file_path = reports_dir / filename
    with open(file_path, "wb") as f:
        f.write(file_stream.getvalue())
    absolute_file_path = str(file_path.resolve())

    # 6. Готовим данные для бота
    bot_internal_base_url = os.getenv("BOT_INTERNAL_URL", "http://bot:8003")
    send_report_endpoint = f"{bot_internal_base_url}/internal/send_write_off_report"
    bot_payload = {
        "chat_id": str(group_id),
        "file_path": absolute_file_path,
        "photos": photos,  # Массив фотографий для отправки
        "photos_count": len(photos),
        "items_count": len(items)
    }

    # 7. Логируем результат и запускаем background-задачу
    logger.info(f"📊 Создан отчет для группы {group_id}: {len(items)} позиций, {len(photos)} фотографий")
    logger.info(f"📄 Документ: {absolute_file_path}")
    background_tasks.add_task(send_write_off_report_to_bot, send_report_endpoint, bot_payload, absolute_file_path)

    return {
        "status": "success",
        "message": f"Запрос на формирование и отправку акта списания получен. Бот скоро отправит {len(photos)} фото и файл в группу.",
        "chat_id": group_id,
        "file_path": absolute_file_path,
        "photos_count": len(photos),
        "items_count": len(items)
    }

async def send_write_off_report_to_bot(url: str, payload: dict, file_path_to_delete: str):
    """
    Отправляет акт списания боту (фото + DOCX) и удаляет временный файл.
    """
    async with httpx.AsyncClient(timeout=60.0) as client:
        try:
            logger.info(f"🤖 Отправляем отчет боту: {payload['photos_count']} фото + DOCX")
            logger.info(f"📡 URL: {url}")
            logger.info(f"💬 Чат: {payload['chat_id']}")
            
            response = await client.post(url, json=payload)
            response.raise_for_status()
            
            logger.info(f"✅ Отчет успешно отправлен боту для чата {payload['chat_id']}")
            
            # Удаляем файл после успешной отправки
            try:
                os.remove(file_path_to_delete)
                logger.info(f"🗑️ Временный файл удален: {file_path_to_delete}")
            except OSError as e:
                logger.warning(f"⚠️ Не удалось удалить файл {file_path_to_delete}: {e}")
        except Exception as e:
            logger.error(f"❌ Ошибка при отправке отчета боту: {e}")
            logger.error(f"📄 Файл остался: {file_path_to_delete}") 