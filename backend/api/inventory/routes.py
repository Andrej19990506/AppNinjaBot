"""
Маршруты для работы с инвентарем
"""
import os
import json
import logging
import traceback
import uuid
from pathlib import Path
from datetime import datetime
from io import BytesIO
from typing import Dict, List, Optional, Union, Any
import pytz
import openpyxl
import pandas as pd
from flask import request, jsonify, current_app, send_file, Response, make_response
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import inventory_bp
from data.history import ItemHistory
from data.write_offs import get_chat_write_offs, add_write_off, update_write_off, delete_write_off

# Логгер
logger = logging.getLogger(__name__)

# Временно используем прямой доступ к путям, в будущем можно переместить
# в конфигурацию или контекст приложения
def get_inventory_path(chat_id):
    """Получить путь к файлу инвентаря"""
    # В будущем эту функцию можно перенести в отдельный модуль utils
    data_dir = current_app.config.get('DATA_DIR')
    inventory_dir = data_dir / 'inventory'
    return inventory_dir / f'inventory_{chat_id}.json'

# Маршрут для работы с инвентарем конкретного чата
@inventory_bp.route('/<chat_id>', methods=['GET'])
def get_chat_inventory(chat_id):
    """Получение инвентаря для конкретного чата"""
    try:
        inventory_path = get_inventory_path(chat_id)
        if not inventory_path.exists():
            logger.warning(f"Inventory file not found for chat {chat_id}")
            return jsonify({
                'inventory': {},
                'metadata': {
                    'lastUpdated': datetime.now().isoformat(),
                    'progress': 0
                }
            })
        
        with open(inventory_path, 'r', encoding='utf-8') as f:
            inventory_data = json.load(f)
            logger.info(f"Successfully loaded inventory for chat {chat_id}")
            return jsonify(inventory_data)
    except Exception as e:
        logger.error(f"Error getting inventory for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/<chat_id>', methods=['POST'])
def update_chat_inventory(chat_id):
    """Обновление инвентаря для конкретного чата"""
    try:
        logger.info(f"Updating inventory for chat {chat_id}")
        data = request.get_json()
        
        if not data:
            return jsonify({'error': 'No data provided'}), 400
        
        inventory_path = get_inventory_path(chat_id)
        logger.info(f"Inventory path: {inventory_path}")
        
        if inventory_path.exists():
            # Загружаем существующий инвентарь
            try:
                with open(inventory_path, 'r', encoding='utf-8') as f:
                    existing_data = json.load(f)
                    logger.info(f"Successfully loaded existing inventory for chat {chat_id}")
            except Exception as e:
                logger.error(f"Error loading existing inventory: {str(e)}")
                existing_data = {'inventory': {}, 'metadata': {}}
        else:
            existing_data = {'inventory': {}, 'metadata': {}}
        
        # Извлекаем инвентарь и метаданные
        if 'inventory' in data:
            inventory = data['inventory']
        else:
            inventory = data  # Считаем, что весь объект - это инвентарь
            data = {'inventory': inventory, 'metadata': existing_data.get('metadata', {})}
        
        # Обновляем метаданные
        metadata = data.get('metadata', {})
        if not metadata.get('lastUpdated'):
            metadata['lastUpdated'] = datetime.now().isoformat()
        
        # Сохраняем историю изменений, если указаны необходимые данные
        if 'history' in data:
            save_to_history(chat_id, data)
        
        # Сохраняем обновленный инвентарь
        os.makedirs(os.path.dirname(inventory_path), exist_ok=True)
        with open(inventory_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
            logger.info(f"Successfully saved inventory for chat {chat_id}")

        # Проверяем, обновлен ли инвентарь, и шлем данные для вебсокетов
        # Если импорт не работает, то отправка будет просто пропущена
        try:
            from ws_module.broadcasters import broadcast_inventory_update
            broadcast_inventory_update(chat_id)
            logger.info(f"Broadcast inventory update for chat {chat_id}")
        except ImportError:
            logger.warning("WebSocket module not available, skipping broadcast")

        return jsonify({'status': 'success', 'timestamp': metadata.get('lastUpdated')})
    except Exception as e:
        logger.error(f"Error updating inventory for chat {chat_id}: {str(e)}")
        logger.error(traceback.format_exc())
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/<chat_id>', methods=['DELETE'])
def delete_chat_inventory(chat_id):
    """Удаление инвентаря для конкретного чата"""
    try:
        inventory_path = get_inventory_path(chat_id)
        
        if not inventory_path.exists():
            return jsonify({'error': 'Inventory not found'}), 404
        
        os.remove(inventory_path)
        logger.info(f"Deleted inventory for chat {chat_id}")
        
        return jsonify({'status': 'success'})
    except Exception as e:
        logger.error(f"Error deleting inventory for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/<chat_id>/status', methods=['GET'])
def get_inventory_status(chat_id):
    """Получение статуса инвентаря для конкретного чата"""
    try:
        inventory_path = get_inventory_path(chat_id)
        
        if not inventory_path.exists():
            return jsonify({
                'exists': False,
                'lastUpdated': None,
                'size': 0,
                'progress': 0
            })
        
        with open(inventory_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        # Проверяем структуру данных
        metadata = data.get('metadata', {})
        if 'inventory' in data:
            inventory = data['inventory']
        else:
            inventory = data
            
        # Подсчитываем общее количество товаров
        total_items = 0
        filled_items = 0
        
        for category, items in inventory.items():
            for item, details in items.items():
                total_items += 1
                if details.get('quantity') is not None and details.get('quantity') != 0:
                    filled_items += 1
        
        # Рассчитываем прогресс
        progress = (filled_items / total_items * 100) if total_items > 0 else 0
        
        # Формируем ответ
        response = {
            'exists': True,
            'lastUpdated': metadata.get('lastUpdated', datetime.now().isoformat()),
            'size': os.path.getsize(inventory_path),
            'progress': progress,
            'total_items': total_items,
            'filled_items': filled_items
        }
        
        return jsonify(response)
    except Exception as e:
        logger.error(f"Error getting inventory status for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/<chat_id>/history', methods=['GET'])
def get_inventory_history(chat_id):
    """Получение истории изменений инвентаря для конкретного чата"""
    try:
        # Получаем экземпляр ItemHistory
        item_history = current_app.item_history
        
        # Получаем историю
        history = item_history.get_chat_history(chat_id)
        
        return jsonify(history)
    except Exception as e:
        logger.error(f"Error getting inventory history for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

def save_to_history(chat_id, inventory_data):
    """Сохранение данных в истории"""
    try:
        if 'history' not in inventory_data:
            return
        
        history_records = inventory_data['history']
        if not history_records:
            return
        
        # Получаем экземпляр ItemHistory
        item_history = current_app.item_history
        
        # Добавляем каждую запись
        for record in history_records:
            try:
                item_history.add_record(chat_id, record)
            except Exception as e:
                logger.error(f"Error adding history record: {str(e)}")
                continue
        
        logger.info(f"Added {len(history_records)} history records for chat {chat_id}")
    except Exception as e:
        logger.error(f"Error saving history for chat {chat_id}: {str(e)}")

@inventory_bp.route('/<chat_id>/excel', methods=['GET', 'POST'])
def generate_excel(chat_id):
    """Генерация Excel-отчета по инвентарю"""
    try:
        if request.method == 'POST':
            # Получаем параметры из запроса
            data = request.get_json()
            categories = data.get('categories', [])
            
            # Если категории не указаны, используем все
            if not categories:
                inventory_path = get_inventory_path(chat_id)
                if not inventory_path.exists():
                    return jsonify({'error': 'Inventory not found'}), 404
                
                with open(inventory_path, 'r', encoding='utf-8') as f:
                    inventory_data = json.load(f)
                
                if 'inventory' in inventory_data:
                    inventory = inventory_data['inventory']
                else:
                    inventory = inventory_data
                
                categories = list(inventory.keys())
            
            # Генерируем временное имя файла
            temp_file = current_app.config.get('TEMPLATES_DIR') / f'export_{chat_id}_{uuid.uuid4()}.xlsx'
            
            # Создаем DataFrame для Excel
            data_frames = []
            
            # Загружаем инвентарь
            inventory_path = get_inventory_path(chat_id)
            with open(inventory_path, 'r', encoding='utf-8') as f:
                inventory_data = json.load(f)
            
            if 'inventory' in inventory_data:
                inventory = inventory_data['inventory']
            else:
                inventory = inventory_data
            
            # Перебираем выбранные категории
            for category in categories:
                if category not in inventory:
                    continue
                
                items = inventory[category]
                rows = []
                
                for item_name, item_data in items.items():
                    quantity = item_data.get('quantity', 0)
                    
                    # Добавляем данные только для товаров с количеством больше 0
                    if quantity > 0:
                        rows.append({
                            'Наименование': item_name,
                            'Количество': quantity,
                            'Ед. изм.': item_data.get('unit', 'шт.')
                        })
                
                # Если есть данные, добавляем в список DataFrame
                if rows:
                    df = pd.DataFrame(rows)
                    df = df.sort_values(by='Наименование')
                    data_frames.append((category, df))
            
            # Создаем Excel-файл
            with pd.ExcelWriter(temp_file, engine='openpyxl') as writer:
                # Для каждой категории создаем отдельный лист
                for category, df in data_frames:
                    sheet_name = category[:31]  # Ограничение Excel на длину имени листа
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    
                    # Получаем созданный лист
                    worksheet = writer.sheets[sheet_name]
                    
                    # Настраиваем ширину столбцов
                    for idx, col in enumerate(df.columns):
                        max_length = max(df[col].astype(str).map(len).max(), len(str(col)))
                        adjusted_width = max_length + 5  # Добавляем отступ
                        worksheet.column_dimensions[get_column_letter(idx + 1)].width = adjusted_width
                    
                    # Настраиваем стиль заголовков
                    for cell in worksheet[1]:
                        cell.font = Font(bold=True)
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                        cell.fill = PatternFill(start_color='A9D08E', end_color='A9D08E', fill_type='solid')
                        cell.border = Border(
                            left=Side(style='thin', color='000000'),
                            right=Side(style='thin', color='000000'),
                            top=Side(style='thin', color='000000'),
                            bottom=Side(style='thin', color='000000')
                        )
            
            # Отправляем файл
            with open(temp_file, 'rb') as f:
                file_data = BytesIO(f.read())
            
            # Удаляем временный файл
            os.remove(temp_file)
            
            # Текущая дата для имени файла
            today = datetime.now().strftime('%Y-%m-%d')
            
            # Отправляем файл клиенту
            return send_file(
                file_data,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True,
                download_name=f'inventory_{chat_id}_{today}.xlsx'
            )
        else:
            # Для GET-запроса возвращаем список категорий
            inventory_path = get_inventory_path(chat_id)
            if not inventory_path.exists():
                return jsonify({'error': 'Inventory not found'}), 404
            
            with open(inventory_path, 'r', encoding='utf-8') as f:
                inventory_data = json.load(f)
            
            if 'inventory' in inventory_data:
                inventory = inventory_data['inventory']
            else:
                inventory = inventory_data
            
            # Собираем статистику по категориям
            categories = []
            for category, items in inventory.items():
                total_items = len(items)
                filled_items = sum(1 for item in items.values() if item.get('quantity', 0) > 0)
                
                categories.append({
                    'name': category,
                    'total_items': total_items,
                    'filled_items': filled_items
                })
            
            return jsonify(categories)
    except Exception as e:
        logger.error(f"Error generating Excel for chat {chat_id}: {str(e)}")
        logger.error(traceback.format_exc())
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/<chat_id>/excel-preview', methods=['GET'])
def preview_excel(chat_id):
    """Предпросмотр данных для Excel-отчета"""
    try:
        # Загружаем инвентарь
        inventory_path = get_inventory_path(chat_id)
        if not inventory_path.exists():
            return jsonify({'error': 'Inventory not found'}), 404
        
        with open(inventory_path, 'r', encoding='utf-8') as f:
            inventory_data = json.load(f)
        
        if 'inventory' in inventory_data:
            inventory = inventory_data['inventory']
        else:
            inventory = inventory_data
        
        # Собираем данные для предпросмотра
        preview_data = {}
        
        for category, items in inventory.items():
            category_items = []
            
            for item_name, item_data in items.items():
                quantity = item_data.get('quantity', 0)
                
                # Добавляем данные только для товаров с количеством больше 0
                if quantity > 0:
                    category_items.append({
                        'name': item_name,
                        'quantity': quantity,
                        'unit': item_data.get('unit', 'шт.')
                    })
            
            # Сортируем по имени
            category_items.sort(key=lambda x: x['name'])
            
            # Если есть товары, добавляем категорию
            if category_items:
                preview_data[category] = category_items
        
        return jsonify(preview_data)
    except Exception as e:
        logger.error(f"Error generating Excel preview for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/templates/inventory_template', methods=['GET'])
def get_inventory_template():
    """Получение шаблона инвентаря"""
    try:
        logger.info('Запрос шаблона инвентаря')
        
        template_path = current_app.config.get('TEMPLATE_PATH')
        if not template_path.exists():
            logger.error(f"Шаблон не найден по пути {template_path}")
            return jsonify({'error': 'Template not found'}), 404
            
        with open(template_path, 'r', encoding='utf-8') as f:
            template = json.load(f)
            logger.info('Шаблон успешно загружен')
            return jsonify(template)
    except Exception as e:
        logger.error(f"Error getting inventory template: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/templates/inventory_template', methods=['POST', 'PATCH', 'PUT'])
def update_inventory_template():
    """Обновление шаблона инвентаря"""
    try:
        logger.info('Обновление шаблона инвентаря')
        data = request.get_json()
        
        if not data or not isinstance(data, dict):
            return jsonify({'error': 'Invalid template data'}), 400
        
        template_path = current_app.config.get('TEMPLATE_PATH')
        
        # Проверяем, есть ли уже шаблон
        if template_path.exists():
            with open(template_path, 'r', encoding='utf-8') as f:
                current_template = json.load(f)
        else:
            current_template = {}
        
        # Для PATCH и PUT обрабатываем по-разному
        if request.method == 'PUT':
            # PUT заменяет весь шаблон
            new_template = data
        else:
            # POST и PATCH обновляют существующий шаблон
            new_template = current_template.copy()
            
            # Обрабатываем добавление новой категории и товаров
            for category, items in data.items():
                if category in new_template:
                    # Обновляем существующую категорию
                    new_template[category].update(items)
                else:
                    # Добавляем новую категорию
                    new_template[category] = items
        
        # Сохраняем обновленный шаблон
        with open(template_path, 'w', encoding='utf-8') as f:
            json.dump(new_template, f, ensure_ascii=False, indent=2)
        
        logger.info('Шаблон успешно обновлен')
        return jsonify({'status': 'success'})
    except Exception as e:
        logger.error(f"Error updating inventory template: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/templates/inventory_template', methods=['DELETE'])
def delete_from_inventory_template():
    """Удаление товара из шаблона инвентаря"""
    try:
        logger.info('=== Удаление товара из шаблона ===')
        data = request.get_json()
        logger.info(f'Полученные данные: {data}')
        
        # Проверяем наличие всех необходимых полей
        required_fields = ['category', 'item']
        if not all(field in data for field in required_fields):
            missing_fields = [field for field in required_fields if field not in data]
            logger.error(f'Отсутствуют обязательные поля: {missing_fields}')
            return jsonify({
                'error': 'Missing required fields',
                'missing_fields': missing_fields
            }), 400
        
        category = data['category']
        item = data['item']
        
        # Загружаем текущий шаблон
        template_path = current_app.config.get('TEMPLATE_PATH')
        logger.info(f'Загружаем шаблон из {template_path}')
        try:
            with open(template_path, 'r', encoding='utf-8') as f:
                template = json.load(f)
            logger.info(f'Текущий шаблон: {template}')
        except Exception as e:
            logger.error(f'Ошибка при загрузке шаблона: {str(e)}')
            return jsonify({'error': 'Failed to load template'}), 500
        
        # Проверяем существование категории и товара
        if category not in template:
            logger.error(f'Категория "{category}" не найдена в шаблоне')
            return jsonify({'error': 'Category not found'}), 404
        
        if item not in template[category]:
            logger.error(f'Товар "{item}" не найден в категории "{category}"')
            return jsonify({'error': 'Item not found'}), 404
        
        # Удаляем товар
        del template[category][item]
        
        # Если категория пуста, удаляем её
        if not template[category]:
            del template[category]
        
        try:
            # Сохраняем обновленный шаблон
            logger.info('Сохраняем обновленный шаблон')
            with open(template_path, 'w', encoding='utf-8') as f:
                json.dump(template, f, ensure_ascii=False, indent=2)
            
            logger.info(f'Товар "{item}" успешно удален из категории "{category}"')
            return jsonify({'status': 'success'})
        except Exception as e:
            logger.error(f'Ошибка при сохранении шаблона: {str(e)}')
            return jsonify({'error': 'Failed to save template'}), 500
    except Exception as e:
        logger.error(f"Error deleting from inventory template: {str(e)}")
        return jsonify({'error': str(e)}), 500

def update_template_after_approval(category: str, item: str) -> bool:
    """Обновление шаблона после подтверждения"""
    try:
        template_path = current_app.config.get('TEMPLATE_PATH')
        
        # Загружаем текущий шаблон
        with open(template_path, 'r', encoding='utf-8') as f:
            template = json.load(f)
        
        # Проверяем наличие категории
        if category not in template:
            template[category] = {}
        
        # Проверяем наличие товара
        if item not in template[category]:
            template[category][item] = {'unit': 'шт.'}
            
            # Сохраняем обновленный шаблон
            with open(template_path, 'w', encoding='utf-8') as f:
                json.dump(template, f, ensure_ascii=False, indent=2)
            
            logger.info(f'Товар "{item}" успешно добавлен в категорию "{category}" шаблона')
            return True
        
        return False  # Товар уже существует
    except Exception as e:
        logger.error(f'Ошибка при обновлении шаблона: {str(e)}')
        return False

@inventory_bp.route('/delete_item', methods=['POST'])
def delete_item():
    """Удаление товара из инвентаря всех чатов"""
    try:
        data = request.get_json()
        
        # Проверяем наличие всех необходимых полей
        required_fields = ['category', 'item']
        if not all(field in data for field in required_fields):
            missing_fields = [field for field in required_fields if field not in data]
            return jsonify({
                'error': 'Missing required fields',
                'missing_fields': missing_fields
            }), 400
        
        category = data['category']
        item = data['item']
        
        # Получаем все файлы инвентаря
        data_dir = current_app.config.get('DATA_DIR')
        inventory_dir = data_dir / 'inventory'
        inventory_files = [f for f in os.listdir(inventory_dir) if f.startswith('inventory_') and f.endswith('.json')]
        
        # Список чатов, где был удален товар
        updated_chats = []
        
        # Обрабатываем каждый файл инвентаря
        for inventory_file in inventory_files:
            chat_id = inventory_file.replace('inventory_', '').replace('.json', '')
            inventory_path = inventory_dir / inventory_file
            
            # Загружаем инвентарь
            with open(inventory_path, 'r', encoding='utf-8') as f:
                inventory_data = json.load(f)
            
            # Проверяем структуру данных
            if 'inventory' in inventory_data:
                inventory = inventory_data['inventory']
            else:
                inventory = inventory_data
                inventory_data = {'inventory': inventory}
            
            # Проверяем наличие категории и товара
            if category in inventory and item in inventory[category]:
                # Удаляем товар
                del inventory[category][item]
                
                # Если категория пуста, удаляем её
                if not inventory[category]:
                    del inventory[category]
                
                # Сохраняем обновленный инвентарь
                with open(inventory_path, 'w', encoding='utf-8') as f:
                    json.dump(inventory_data, f, ensure_ascii=False, indent=2)
                
                updated_chats.append(chat_id)
                
                # Отправляем уведомление через вебсокет
                try:
                    from ws_module.broadcasters import broadcast_inventory_update
                    broadcast_inventory_update(chat_id)
                except ImportError:
                    logger.warning("WebSocket module not available, skipping broadcast")
        
        # Удаляем товар из шаблона
        template_path = current_app.config.get('TEMPLATE_PATH')
        with open(template_path, 'r', encoding='utf-8') as f:
            template = json.load(f)
        
        if category in template and item in template[category]:
            del template[category][item]
            
            # Если категория пуста, удаляем её
            if not template[category]:
                del template[category]
            
            with open(template_path, 'w', encoding='utf-8') as f:
                json.dump(template, f, ensure_ascii=False, indent=2)
        
        return jsonify({
            'status': 'success',
            'updated_chats': updated_chats
        })
    except Exception as e:
        logger.error(f"Error deleting item: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/item_history/<chat_id>', methods=['POST'])
def add_history_record(chat_id):
    """Добавление записи в историю инвентаря"""
    try:
        data = request.get_json()
        
        if not data:
            return jsonify({'error': 'No data provided'}), 400
        
        # Получаем экземпляр ItemHistory
        item_history = current_app.item_history
        
        # Добавляем запись в историю
        result = item_history.add_record(chat_id, data)
        
        return jsonify(result)
    except Exception as e:
        logger.error(f"Error adding history record for chat {chat_id}: {str(e)}")
        return jsonify({'error': str(e)}), 500

@inventory_bp.route('/item_history/<chat_id>/<category>/<path:item>', methods=['GET'])
def get_item_history_endpoint(chat_id, category, item):
    """Получение истории изменений конкретного товара"""
    try:
        # Получаем экземпляр ItemHistory
        item_history = current_app.item_history
        
        # Получаем историю товара
        history = item_history.get_item_history(chat_id, category, item)
        
        return jsonify(history)
    except Exception as e:
        logger.error(f"Error getting item history for chat {chat_id}, category {category}, item {item}: {str(e)}")
        return jsonify({'error': str(e)}), 500 