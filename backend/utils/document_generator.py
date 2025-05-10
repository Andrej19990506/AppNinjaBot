"""
Модуль для генерации документов DOCX с актами списания
"""
import io
import os
from datetime import datetime
from typing import Dict, List, Optional
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_write_off_document(data: Dict) -> io.BytesIO:
    """
    Создает DOCX документ акта списания на основе переданных данных
    
    Args:
        data: словарь с данными для акта списания
            - chatTitle: название чата/филиала
            - date: дата акта списания
            - items: список элементов списания
    
    Returns:
        Документ в формате BytesIO
    """
    # Создаем документ
    doc = Document()
    
    # Устанавливаем поля страницы
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
    
    # Добавляем заголовок документа
    title = doc.add_paragraph("Акт списания")
    title_format = title.paragraph_format
    title_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_format.space_after = Pt(12)
    title_run = title.runs[0]
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    
    # Форматирование даты и филиала
    chat_title = data.get('chatTitle', 'Неизвестный филиал')
    
    # Форматируем дату из ISO в читаемый формат
    date_str = data.get('date', datetime.now().isoformat())
    try:
        if isinstance(date_str, str):
            date_obj = datetime.fromisoformat(date_str.replace('Z', '+00:00'))
            date_formatted = date_obj.strftime("%d.%m.%Y %H:%M")
        else:
            date_formatted = datetime.now().strftime("%d.%m.%Y %H:%M")
    except (ValueError, TypeError):
        date_formatted = datetime.now().strftime("%d.%m.%Y %H:%M")
    
    # Добавляем строку с датой и филиалом
    date_line = doc.add_paragraph(f"От {date_formatted}. Филиал: {chat_title}.")
    date_line.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line.paragraph_format.space_after = Pt(12)
    
    # Создаем таблицу для элементов списания
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Устанавливаем ширину столбцов
    table.columns[0].width = Cm(6)  # Наименование
    table.columns[1].width = Cm(3)  # Кол-во
    table.columns[2].width = Cm(6)  # Причина списания
    
    # Заголовки таблицы
    header_cells = table.rows[0].cells
    header_cells[0].text = "Наименование"
    header_cells[1].text = "Кол-во"
    header_cells[2].text = "Причина списания"
    
    # Стиль заголовка таблицы
    for cell in header_cells:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    
    # Добавляем строки с элементами списания
    items = data.get('items', [])
    for item in items:
        row_cells = table.add_row().cells
        
        # Наименование
        row_cells[0].text = item.get('name', '')
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        # Количество с единицей измерения
        quantity = item.get('quantity', 0)
        unit_type = item.get('unitType', 'шт')
        row_cells[1].text = f"{quantity} {unit_type}"
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Причина списания
        reason = item.get('reason', {})
        reason_title = reason.get('title', '') if reason else ''
        
        # Добавляем детальное описание в скобках, если оно существует
        description = item.get('description', '')
        if description:
            reason_text = f"{reason_title} ({description})"
        else:
            reason_text = reason_title
            
        row_cells[2].text = reason_text
        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    # Добавляем информацию о времени формирования документа
    doc.add_paragraph("")
    generation_info = doc.add_paragraph(f"Документ сформирован: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}")
    generation_info.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    generation_info.paragraph_format.space_after = Pt(12)
    
    # Добавляем раздел "Пояснения:"
    explanations = doc.add_paragraph("Пояснения:")
    explanations.runs[0].font.bold = True
    explanations.paragraph_format.space_after = Pt(6)
    
    # Добавляем пояснения к списанию
    doc.add_paragraph("Расход: замена масла во фритюре, списание теста на чистку пресса, расход муки;")
    doc.add_paragraph("Отказ клиента: Клиент отказался, не открыл, отменили заказ, изменили готовый заказ и т.п., косяк программы с заказом;")
    doc.add_paragraph("Брак: списывается только пицца (другая продукция на брак не списывается). Разрешено списывать 1-2 шт в день.")
    doc.add_paragraph("Маркетинг: заказ пицц и прочего для съемки, блоггеру, для инстаграмма и т.п., пиццы для мероприятий;")
    doc.add_paragraph("Удержание из з/п: косячно сделан заказ по вине сотрудников, испорчены продукты по вине сотрудников.")
    doc.add_paragraph("Проработка: блюда или использование отдельных продуктов для проработки.")
    doc.add_paragraph("Порча: : порча продуктов (плесень, сок, гниль), истечение срока годности, лом, бой (макаруны, яйца), брак поставщика (мятые коробки, треснутые контейнеры ит.д.).")
    
    # Добавляем раздел с подписями
    doc.add_paragraph("")
    doc.add_paragraph("")
    signatures = doc.add_paragraph("Ответственный: ________________________ / _________________")
    signatures.paragraph_format.space_after = Pt(12)
    
    doc.add_paragraph("Администратор: ________________________ / _________________")
    doc.add_paragraph("")
    doc.add_paragraph("Руководитель: ________________________ / _________________")
    
    # Сохраняем документ в памяти
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    
    return file_stream 